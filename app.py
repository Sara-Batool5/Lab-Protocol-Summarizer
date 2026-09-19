"""
Lab Protocol & Microbiology Summarizer
--------------------------------------
Upload up to 5 research papers / lab protocols (PDF or Google Drive link) and get
structured experimental procedures, reagent lists and safety precautions.

Stack : Streamlit + Groq API (openai/gpt-oss-120b) + SQLAlchemy storage
Secrets (Streamlit Cloud -> App settings -> Secrets):
    GROQ_API_KEY = "gsk_..."          # required
    DATABASE_URL = "postgresql://..." # optional, recommended for real 7-day persistence
"""

from __future__ import annotations

import datetime as dt
import hashlib
import html
import io
import json
import os
import re
import time
import uuid
from pathlib import Path

import pandas as pd
import requests
import streamlit as st
from docx import Document
from docx.shared import Pt
from groq import Groq, RateLimitError
from pypdf import PdfReader
from sqlalchemy import create_engine, text

# --------------------------------------------------------------------------- #
# Page setup
# --------------------------------------------------------------------------- #
st.set_page_config(
    page_title="Lab Protocol & Microbiology Summarizer",
    page_icon="🧫",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --------------------------------------------------------------------------- #
# Settings
# --------------------------------------------------------------------------- #
MODEL = "openai/gpt-oss-120b"
REASONING_EFFORT = "low"          # "low" | "medium" | "high"
MAX_PAPERS = 5
RETENTION_DAYS = 7
CHUNK_CHARS = 14000               # ~3.5k tokens per call. Raise on a paid Groq tier.
MAX_CHUNKS_PER_PAPER = 8          # very long papers: the most method-heavy sections are used
MAX_OUTPUT_TOKENS = 4000
MAX_PDF_MB = 50
DATA_DIR = Path("data")

# --------------------------------------------------------------------------- #
# Styling
# --------------------------------------------------------------------------- #
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;500;600&family=Source+Serif+4:wght@600;700&display=swap');

html, body, .stApp, [data-testid="stMarkdownContainer"], button, input, textarea {
    font-family: 'IBM Plex Sans', sans-serif;
}
h1, h2, h3, .lab-title {
    font-family: 'Source Serif 4', Georgia, serif !important;
    letter-spacing: -0.01em;
}
h2, h3 { color: #0B4F4A; }
.block-container { padding-top: 2rem; max-width: 1150px; }
#MainMenu, footer { visibility: hidden; }

[data-testid="stSidebar"] { background: #EAF4F0; border-right: 1px solid #CFE3DC; }

/* Hero: a petri-dish ring motif in the corner */
.lab-hero {
    position: relative; overflow: hidden;
    background: #0B4F4A; color: #FFFFFF;
    border-radius: 14px; padding: 1.9rem 2.2rem; margin-bottom: 1rem;
}
.lab-hero::after {
    content: ""; position: absolute; right: -80px; top: -80px;
    width: 320px; height: 320px; border-radius: 50%;
    background: radial-gradient(circle,
        transparent 0 44%, rgba(255,255,255,.14) 44% 45%,
        transparent 45% 58%, rgba(255,255,255,.10) 58% 59%,
        transparent 59% 72%, rgba(255,255,255,.07) 72% 73%,
        transparent 73%);
}
.lab-title { font-size: 2.05rem; font-weight: 700; line-height: 1.15; margin: 0 0 .35rem 0; color: #FFFFFF; }
.lab-sub { font-size: 1.02rem; color: #CDEBE3; max-width: 46rem; margin: 0; }

.lab-notice {
    background: #FFF7E3; border-left: 4px solid #B7791F; color: #4A3510;
    border-radius: 8px; padding: .7rem 1rem; margin-bottom: 1.1rem; font-size: .95rem;
}
.lab-side-note { font-size: .85rem; color: #35524E; line-height: 1.45; }
.chip {
    display: inline-block; padding: 2px 11px; border-radius: 999px;
    background: #DDF0EA; color: #0B4F4A; font-size: .8rem; margin: 0 6px 6px 0;
}
.chip.warn { background: #FBEBC8; color: #6B4A0B; }
.lab-empty {
    border: 1px dashed #9CC7BC; border-radius: 12px; padding: 1.6rem 1.8rem;
    background: #F6FAF8; color: #23413D;
}
</style>
"""

# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #
def _secret(name: str, default: str | None = None) -> str | None:
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass
    return os.environ.get(name, default)


def _stretch(fn, *args, **kwargs):
    """Call a Streamlit widget at full width across Streamlit versions."""
    try:
        return fn(*args, width="stretch", **kwargs)
    except Exception:
        return fn(*args, use_container_width=True, **kwargs)


def now_utc() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)


def iso(t: dt.datetime) -> str:
    return t.isoformat(timespec="seconds")


def fmt_date(s: str) -> str:
    try:
        return dt.datetime.fromisoformat(s).strftime("%d %b %Y")
    except Exception:
        return s


def md_escape(s: str) -> str:
    return re.sub(r"([\\`*_{}\[\]<>#~|$])", r"\\\1", s or "")


def safe_name(name: str) -> str:
    stem = Path(name).stem
    return re.sub(r"[^\w\-]+", "_", stem).strip("_")[:60] or "paper"


# --------------------------------------------------------------------------- #
# Storage (SQLite by default, Postgres if DATABASE_URL is set)
# --------------------------------------------------------------------------- #
DDL = """
CREATE TABLE IF NOT EXISTS results (
    id TEXT PRIMARY KEY,
    workspace TEXT NOT NULL,
    filename TEXT,
    source TEXT,
    text_hash TEXT,
    created_at TEXT,
    expires_at TEXT,
    result_json TEXT
)
"""


@st.cache_resource(show_spinner=False)
def get_engine():
    url = _secret("DATABASE_URL")
    if url:
        if url.startswith("postgres://"):
            url = url.replace("postgres://", "postgresql+psycopg2://", 1)
        elif url.startswith("postgresql://"):
            url = url.replace("postgresql://", "postgresql+psycopg2://", 1)
        engine = create_engine(url, pool_pre_ping=True, pool_recycle=300)
    else:
        DATA_DIR.mkdir(exist_ok=True)
        engine = create_engine(
            f"sqlite:///{DATA_DIR / 'labsum.db'}",
            connect_args={"check_same_thread": False},
        )
    with engine.begin() as conn:
        conn.execute(text(DDL))
    return engine


def purge_expired() -> None:
    with get_engine().begin() as conn:
        conn.execute(text("DELETE FROM results WHERE expires_at < :now"), {"now": iso(now_utc())})


def save_result(ws: str, filename: str, source: str, text_hash: str, data: dict) -> str:
    rid = uuid.uuid4().hex
    created = now_utc()
    with get_engine().begin() as conn:
        conn.execute(
            text(
                "INSERT INTO results (id, workspace, filename, source, text_hash, created_at, expires_at, result_json) "
                "VALUES (:id, :ws, :fn, :src, :h, :c, :e, :j)"
            ),
            {
                "id": rid, "ws": ws, "fn": filename, "src": source, "h": text_hash,
                "c": iso(created), "e": iso(created + dt.timedelta(days=RETENTION_DAYS)),
                "j": json.dumps(data, ensure_ascii=False),
            },
        )
    return rid


def find_cached(ws: str, text_hash: str) -> str | None:
    with get_engine().connect() as conn:
        row = conn.execute(
            text("SELECT id FROM results WHERE workspace = :ws AND text_hash = :h AND expires_at >= :now LIMIT 1"),
            {"ws": ws, "h": text_hash, "now": iso(now_utc())},
        ).fetchone()
    return row[0] if row else None


def list_results(ws: str) -> list[dict]:
    with get_engine().connect() as conn:
        rows = conn.execute(
            text(
                "SELECT id, filename, source, created_at, expires_at, result_json FROM results "
                "WHERE workspace = :ws AND expires_at >= :now ORDER BY created_at DESC"
            ),
            {"ws": ws, "now": iso(now_utc())},
        ).fetchall()
    out = []
    for r in rows:
        m = r._mapping
        out.append(
            {
                "id": m["id"], "filename": m["filename"], "source": m["source"],
                "created_at": m["created_at"], "expires_at": m["expires_at"],
                "data": normalize(json.loads(m["result_json"])),
            }
        )
    return out


def delete_result(ws: str, rid: str) -> None:
    with get_engine().begin() as conn:
        conn.execute(text("DELETE FROM results WHERE id = :id AND workspace = :ws"), {"id": rid, "ws": ws})


def get_workspace() -> str:
    ws = st.query_params.get("ws")
    if not ws or not re.fullmatch(r"[a-f0-9]{16}", str(ws)):
        ws = uuid.uuid4().hex[:16]
        st.query_params["ws"] = ws
    return str(ws)


# --------------------------------------------------------------------------- #
# Input: PDF / Google Drive
# --------------------------------------------------------------------------- #
def extract_drive_id(url: str) -> str | None:
    for pat in (r"/file/d/([A-Za-z0-9_-]{10,})", r"[?&]id=([A-Za-z0-9_-]{10,})", r"/d/([A-Za-z0-9_-]{10,})"):
        m = re.search(pat, url)
        if m:
            return m.group(1)
    return None


def fetch_drive_pdf(url: str) -> tuple[str, bytes]:
    fid = extract_drive_id(url)
    if not fid:
        raise ValueError("This does not look like a Google Drive file link. Folder links are not supported.")
    resp = requests.get(
        "https://drive.usercontent.google.com/download",
        params={"id": fid, "export": "download", "confirm": "t"},
        timeout=90,
    )
    if resp.status_code != 200:
        raise ValueError(f"Google Drive returned status {resp.status_code}. Check that the link is shared as 'Anyone with the link'.")
    content = resp.content
    if len(content) > MAX_PDF_MB * 1024 * 1024:
        raise ValueError(f"The file is larger than {MAX_PDF_MB} MB.")
    if not content.lstrip()[:5].startswith(b"%PDF"):
        raise ValueError("Could not download a PDF. Set sharing to 'Anyone with the link can view' and use a direct file link.")
    cd = resp.headers.get("content-disposition", "")
    m = re.search(r'filename="([^"]+)"', cd)
    name = m.group(1) if m else f"drive_{fid[:8]}.pdf"
    return name, content


def pdf_to_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def clean_text(t: str) -> str:
    t = t.replace("\x00", " ")
    t = re.sub(r"(\w)-\n(\w)", r"\1\2", t)
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    matches = list(re.finditer(r"\n\s*(references|bibliography|literature cited)\s*\n", t, flags=re.I))
    if matches and matches[-1].start() > len(t) * 0.5:
        t = t[: matches[-1].start()]
    return t.strip()


def split_chunks(t: str, size: int) -> list[str]:
    chunks, cur, n = [], [], 0
    for line in t.split("\n"):
        if cur and n + len(line) + 1 > size:
            chunks.append("\n".join(cur))
            cur, n = [], 0
        while len(line) > size:
            chunks.append(line[:size])
            line = line[size:]
        cur.append(line)
        n += len(line) + 1
    if cur:
        chunks.append("\n".join(cur))
    return [c for c in chunks if c.strip()]


KEYWORDS = (
    "incubat", "centrifug", "buffer", "medium", "media", "reagent", "protocol", "µl", "μl", " ml",
    "°c", "rpm", "pcr", "culture", "plate", "stain", "dilut", "autoclave", "biosafety", "hazard",
    "safety", "purchased", "sigma", "thermo", "wash", "extract",
)


def pick_chunks(chunks: list[str], k: int) -> tuple[list[str], str]:
    scores = [(sum(c.lower().count(w) for w in KEYWORDS), i) for i, c in enumerate(chunks)]
    keep = {0}
    for _, i in sorted(scores, reverse=True):
        if len(keep) >= k:
            break
        keep.add(i)
    idx = sorted(keep)
    note = f"Long document: the {len(idx)} most method-heavy sections out of {len(chunks)} were analysed."
    return [chunks[i] for i in idx], note


# --------------------------------------------------------------------------- #
# LLM extraction (Groq)
# --------------------------------------------------------------------------- #
SYSTEM_PROMPT = (
    "You are a meticulous life-science research assistant specialised in microbiology, molecular biology "
    "and laboratory protocols. You extract structured information from papers exactly as written. "
    "You never invent reagents, quantities, steps or safety claims. You reply with a single valid JSON object only."
)

USER_PROMPT = """Extract structured information from this portion ({part} of {total}) of a life-science research paper or lab protocol.

Return ONE JSON object with exactly these keys:
{{
  "title": "paper title if visible in this portion, else \\"\\"",
  "study_type": "e.g. in vitro assay, clinical isolate study, protocol, review; else \\"\\"",
  "summary": "3-5 sentence plain-language summary of this portion (objective, approach, key finding)",
  "organisms_and_samples": ["organisms, strains, cell lines, sample types"],
  "procedures": [{{"name": "procedure name", "steps": ["one action per step, in original order"]}}],
  "reagents": [{{"name": "", "concentration_or_amount": "", "supplier_or_catalog": "", "purpose": ""}}],
  "equipment": ["instruments and consumables"],
  "safety_precautions": [{{"hazard": "", "precaution": "", "basis": "Stated in paper" or "Suggested - verify with SDS/biosafety officer"}}],
  "biosafety_level": "BSL-1/2/3 if stated, else \\"\\"",
  "critical_parameters": ["temperatures, times, speeds, pH, concentrations that are critical"]
}}

Rules:
- Use only what is explicitly in the text. If something is not present in this portion, use "" or [].
- Keep every quantity, temperature, time, speed, pH, concentration and unit exactly as written.
- Put reagents in the reagents list (chemicals, media, buffers, antibodies, enzymes, kits, strains as materials). Do not list the same reagent twice.
- safety_precautions: include precautions explicitly stated (basis "Stated in paper"). If hazardous agents (pathogens, toxic or flammable chemicals, radioactive material, sharps, UV) are used without any precaution stated, you may add a few standard, specific precautions with basis "Suggested - verify with SDS/biosafety officer".

TEXT:
\"\"\"
{chunk}
\"\"\"
"""


def get_client() -> Groq | None:
    key = _secret("GROQ_API_KEY")
    return Groq(api_key=key, timeout=120.0, max_retries=0) if key else None


def call_groq(client: Groq, user_prompt: str) -> str:
    kwargs = dict(
        model=MODEL,
        messages=[{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": user_prompt}],
        temperature=0.1,
        max_completion_tokens=MAX_OUTPUT_TOKENS,
        response_format={"type": "json_object"},
    )
    last: Exception | None = None
    for attempt in range(1, 6):
        try:
            try:
                resp = client.chat.completions.create(reasoning_effort=REASONING_EFFORT, **kwargs)
            except TypeError:
                resp = client.chat.completions.create(**kwargs)
            return resp.choices[0].message.content or ""
        except RateLimitError as e:
            last = e
            time.sleep(min(15 * attempt, 60))
        except Exception as e:
            last = e
            if attempt >= 3:
                break
            time.sleep(2 * attempt)
    raise RuntimeError(f"The AI service could not complete the request ({type(last).__name__}). Please try again in a minute.")


def parse_json(raw: str) -> dict:
    raw = raw.strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.M).strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", raw, flags=re.S)
        if not m:
            raise
        return json.loads(m.group(0))


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _list(v) -> list:
    if v is None or v == "":
        return []
    return v if isinstance(v, list) else [v]


def normalize(d) -> dict:
    d = d if isinstance(d, dict) else {}
    procedures = []
    for p in _list(d.get("procedures")):
        if isinstance(p, dict):
            steps = [_s(s.get("step") if isinstance(s, dict) else s) for s in _list(p.get("steps"))]
            steps = [s for s in steps if s]
            if steps:
                procedures.append({"name": _s(p.get("name")) or "Procedure", "steps": steps})
        elif _s(p):
            procedures.append({"name": "Procedure", "steps": [_s(p)]})

    reagents = []
    for r in _list(d.get("reagents")):
        if isinstance(r, dict):
            item = {
                "name": _s(r.get("name")),
                "concentration_or_amount": _s(r.get("concentration_or_amount")),
                "supplier_or_catalog": _s(r.get("supplier_or_catalog")),
                "purpose": _s(r.get("purpose")),
            }
        else:
            item = {"name": _s(r), "concentration_or_amount": "", "supplier_or_catalog": "", "purpose": ""}
        if item["name"]:
            reagents.append(item)

    safety = []
    for s in _list(d.get("safety_precautions")):
        if isinstance(s, dict):
            item = {"hazard": _s(s.get("hazard")), "precaution": _s(s.get("precaution")), "basis": _s(s.get("basis")) or "Not specified"}
        else:
            item = {"hazard": "General", "precaution": _s(s), "basis": "Not specified"}
        if item["precaution"] or item["hazard"]:
            safety.append(item)

    return {
        "title": _s(d.get("title")),
        "study_type": _s(d.get("study_type")),
        "summary": _s(d.get("summary")),
        "organisms_and_samples": [_s(x) for x in _list(d.get("organisms_and_samples")) if _s(x)],
        "procedures": procedures,
        "reagents": reagents,
        "equipment": [_s(x) for x in _list(d.get("equipment")) if _s(x)],
        "safety_precautions": safety,
        "biosafety_level": _s(d.get("biosafety_level")),
        "critical_parameters": [_s(x) for x in _list(d.get("critical_parameters")) if _s(x)],
        "processing_note": _s(d.get("processing_note")),
    }


def _dedupe_strings(items: list[str]) -> list[str]:
    seen, out = set(), []
    for x in items:
        k = x.lower()
        if k not in seen:
            seen.add(k)
            out.append(x)
    return out


def merge_partials(parts: list[dict]) -> dict:
    merged = normalize({})
    for p in parts:
        for key in ("title", "study_type", "summary", "biosafety_level"):
            if not merged[key] and p[key]:
                merged[key] = p[key]
        for key in ("organisms_and_samples", "equipment", "critical_parameters"):
            merged[key] = _dedupe_strings(merged[key] + p[key])

        for proc in p["procedures"]:
            match = next((m for m in merged["procedures"] if m["name"].lower() == proc["name"].lower()), None)
            if match:
                have = {s.lower() for s in match["steps"]}
                match["steps"] += [s for s in proc["steps"] if s.lower() not in have]
            else:
                merged["procedures"].append({"name": proc["name"], "steps": list(proc["steps"])})

        for r in p["reagents"]:
            match = next((m for m in merged["reagents"] if m["name"].lower() == r["name"].lower()), None)
            if match:
                for k, v in r.items():
                    if not match[k] and v:
                        match[k] = v
            else:
                merged["reagents"].append(dict(r))

        have = {(s["hazard"].lower(), s["precaution"].lower()) for s in merged["safety_precautions"]}
        for s in p["safety_precautions"]:
            key = (s["hazard"].lower(), s["precaution"].lower())
            if key not in have:
                have.add(key)
                merged["safety_precautions"].append(dict(s))
    return merged


def process_paper(client: Groq, ws: str, name: str, source: str, raw: bytes, log) -> str:
    log(f"Reading {name}")
    body = clean_text(pdf_to_text(raw))
    if len(body) < 400:
        raise ValueError("No readable text found. The PDF may be a scanned image without a text layer.")

    text_hash = hashlib.sha256(body.encode("utf-8", "ignore")).hexdigest()[:32]
    cached = find_cached(ws, text_hash)
    if cached:
        log(f"{name} was already analysed. Using the saved result.")
        return cached

    chunks = split_chunks(body, CHUNK_CHARS)
    note = ""
    if len(chunks) > MAX_CHUNKS_PER_PAPER:
        chunks, note = pick_chunks(chunks, MAX_CHUNKS_PER_PAPER)

    parts, failed = [], 0
    for i, chunk in enumerate(chunks, 1):
        log(f"Extracting section {i} of {len(chunks)} from {name}")
        try:
            raw_json = call_groq(client, USER_PROMPT.format(part=i, total=len(chunks), chunk=chunk))
            parts.append(normalize(parse_json(raw_json)))
        except (json.JSONDecodeError, ValueError):
            failed += 1
    if not parts:
        raise RuntimeError("The AI response could not be read. Please try again.")
    if failed:
        note = (note + " " if note else "") + f"{failed} section(s) could not be processed and were skipped."

    merged = merge_partials(parts)
    merged["processing_note"] = note
    return save_result(ws, name, source, text_hash, merged)


# --------------------------------------------------------------------------- #
# Exports
# --------------------------------------------------------------------------- #
DISCLAIMER = (
    "AI-generated extraction. Verify every quantity, step and safety measure against the original paper, "
    "the supplier's Safety Data Sheets and your institution's biosafety rules before working in the lab."
)


def reagents_df(d: dict) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Reagent": r["name"],
                "Amount / concentration": r["concentration_or_amount"],
                "Supplier / catalog no.": r["supplier_or_catalog"],
                "Purpose": r["purpose"],
            }
            for r in d["reagents"]
        ],
        columns=["Reagent", "Amount / concentration", "Supplier / catalog no.", "Purpose"],
    )


def safety_df(d: dict) -> pd.DataFrame:
    return pd.DataFrame(
        [{"Hazard": s["hazard"], "Precaution": s["precaution"], "Basis": s["basis"]} for s in d["safety_precautions"]],
        columns=["Hazard", "Precaution", "Basis"],
    )


def to_markdown(d: dict, filename: str) -> str:
    L = [f"# {d['title'] or filename}", "", f"*Source file: {filename}*", ""]
    if d["study_type"]:
        L += [f"**Study type:** {d['study_type']}", ""]
    if d["summary"]:
        L += ["## Summary", "", d["summary"], ""]
    if d["organisms_and_samples"]:
        L += ["## Organisms and samples", ""] + [f"- {x}" for x in d["organisms_and_samples"]] + [""]
    L += ["## Experimental procedures", ""]
    if d["procedures"]:
        for p in d["procedures"]:
            L += [f"### {p['name']}", ""] + [f"{i}. {s}" for i, s in enumerate(p["steps"], 1)] + [""]
    else:
        L += ["No procedures were found.", ""]
    L += ["## Reagents", ""]
    if d["reagents"]:
        L += ["| Reagent | Amount / concentration | Supplier / catalog no. | Purpose |", "|---|---|---|---|"]
        for r in d["reagents"]:
            cells = [r["name"], r["concentration_or_amount"], r["supplier_or_catalog"], r["purpose"]]
            L.append("| " + " | ".join(c.replace("|", "/") or "-" for c in cells) + " |")
        L.append("")
    else:
        L += ["No reagents were found.", ""]
    if d["equipment"]:
        L += ["## Equipment", ""] + [f"- {x}" for x in d["equipment"]] + [""]
    if d["critical_parameters"]:
        L += ["## Critical parameters", ""] + [f"- {x}" for x in d["critical_parameters"]] + [""]
    L += ["## Safety precautions", ""]
    if d["biosafety_level"]:
        L += [f"**Biosafety level:** {d['biosafety_level']}", ""]
    if d["safety_precautions"]:
        for s in d["safety_precautions"]:
            L.append(f"- **{s['hazard'] or 'General'}:** {s['precaution']} *({s['basis']})*")
        L.append("")
    else:
        L += ["No safety precautions were found.", ""]
    L += ["---", f"*{DISCLAIMER}*", ""]
    return "\n".join(L)


def to_docx(d: dict, filename: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(d["title"] or filename, 0)
    doc.add_paragraph(f"Source file: {filename}")
    if d["study_type"]:
        doc.add_paragraph(f"Study type: {d['study_type']}")
    if d["summary"]:
        doc.add_heading("Summary", 1)
        doc.add_paragraph(d["summary"])
    if d["organisms_and_samples"]:
        doc.add_heading("Organisms and samples", 1)
        for x in d["organisms_and_samples"]:
            doc.add_paragraph(x, style="List Bullet")

    doc.add_heading("Experimental procedures", 1)
    if d["procedures"]:
        for p in d["procedures"]:
            doc.add_heading(p["name"], 2)
            for i, s in enumerate(p["steps"], 1):
                doc.add_paragraph(f"{i}. {s}")
    else:
        doc.add_paragraph("No procedures were found.")

    doc.add_heading("Reagents", 1)
    if d["reagents"]:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ["Reagent", "Amount / concentration", "Supplier / catalog no.", "Purpose"]):
            cell.text = label
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for r in d["reagents"]:
            row = table.add_row().cells
            row[0].text, row[1].text = r["name"], r["concentration_or_amount"] or "-"
            row[2].text, row[3].text = r["supplier_or_catalog"] or "-", r["purpose"] or "-"
    else:
        doc.add_paragraph("No reagents were found.")

    if d["equipment"]:
        doc.add_heading("Equipment", 1)
        for x in d["equipment"]:
            doc.add_paragraph(x, style="List Bullet")
    if d["critical_parameters"]:
        doc.add_heading("Critical parameters", 1)
        for x in d["critical_parameters"]:
            doc.add_paragraph(x, style="List Bullet")

    doc.add_heading("Safety precautions", 1)
    if d["biosafety_level"]:
        doc.add_paragraph(f"Biosafety level: {d['biosafety_level']}")
    if d["safety_precautions"]:
        for s in d["safety_precautions"]:
            doc.add_paragraph(f"{s['hazard'] or 'General'}: {s['precaution']} ({s['basis']})", style="List Bullet")
    else:
        doc.add_paragraph("No safety precautions were found.")

    p = doc.add_paragraph()
    run = p.add_run(DISCLAIMER)
    run.italic = True
    run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# UI pieces
# --------------------------------------------------------------------------- #
def render_header() -> None:
    st.markdown(
        """
        <div class="lab-hero">
            <div class="lab-title">Lab Protocol &amp; Microbiology Summarizer</div>
            <p class="lab-sub">Turn long research papers and lab protocols into clear procedures,
            reagent lists and safety precautions you can export.</p>
        </div>
        <div class="lab-notice">
            <b>Your results are saved for 7 days.</b> Each analysis is deleted automatically 7 days after it was created.
            Bookmark this page's link to come back to your saved results.
        </div>
        """,
        unsafe_allow_html=True,
    )


def parse_links(raw: str) -> list[str]:
    seen, out = set(), []
    for tok in re.split(r"[\s,]+", raw or ""):
        if tok.startswith("http") and tok not in seen:
            seen.add(tok)
            out.append(tok)
    return out


def render_sidebar():
    with st.sidebar:
        st.markdown("## Add papers")
        st.caption(f"Up to {MAX_PAPERS} papers per analysis.")
        files = st.file_uploader("Upload PDF files", type=["pdf"], accept_multiple_files=True)
        links_raw = st.text_area(
            "Or paste Google Drive links",
            placeholder="One link per line",
            height=110,
            help="Each file must be shared as 'Anyone with the link can view'.",
        )
        links = parse_links(links_raw)
        total = len(files or []) + len(links)

        if total > MAX_PAPERS:
            st.error(f"{total} papers selected. Remove {total - MAX_PAPERS} to continue.")
        elif total:
            st.caption(f"{total} of {MAX_PAPERS} papers selected.")

        analyze = _stretch(
            st.button, "Analyze papers", type="primary", disabled=(total == 0 or total > MAX_PAPERS)
        )
        st.divider()
        st.markdown(
            f"""<div class="lab-side-note">
            <b>Saved for {RETENTION_DAYS} days.</b> Results stay available here and are then deleted automatically.<br><br>
            Anyone who has this page's link can open your saved results, so avoid sharing it.
            </div>""",
            unsafe_allow_html=True,
        )
    return files or [], links, analyze


def render_result(rec: dict, key: str) -> None:
    d = rec["data"]
    st.markdown(f"### {md_escape(d['title'] or rec['filename'])}")
    chips = (
        f'<span class="chip">File: {html.escape(rec["filename"])}</span>'
        f'<span class="chip">Source: {html.escape(rec["source"] or "")}</span>'
        f'<span class="chip warn">Available until {fmt_date(rec["expires_at"])}</span>'
    )
    if d["biosafety_level"]:
        chips += f'<span class="chip warn">{html.escape(d["biosafety_level"])}</span>'
    st.markdown(chips, unsafe_allow_html=True)

    t_over, t_proc, t_reag, t_safe, t_exp = st.tabs(
        ["Overview", "Procedures", "Reagents and equipment", "Safety", "Export"]
    )

    with t_over:
        if d["study_type"]:
            st.markdown(f"**Study type:** {md_escape(d['study_type'])}")
        st.write(d["summary"] or "No summary available.")
        if d["organisms_and_samples"]:
            st.markdown("**Organisms and samples**")
            st.markdown("\n".join(f"- {md_escape(x)}" for x in d["organisms_and_samples"]))
        if d["critical_parameters"]:
            st.markdown("**Critical parameters**")
            st.markdown("\n".join(f"- {md_escape(x)}" for x in d["critical_parameters"]))
        if d["processing_note"]:
            st.caption(d["processing_note"])

    with t_proc:
        if not d["procedures"]:
            st.info("No experimental procedures were found in this document.")
        for p in d["procedures"]:
            with st.expander(f"{p['name']} ({len(p['steps'])} steps)", expanded=len(d["procedures"]) <= 3):
                st.markdown("\n".join(f"{i}. {md_escape(s)}" for i, s in enumerate(p["steps"], 1)))

    with t_reag:
        if d["reagents"]:
            _stretch(st.dataframe, reagents_df(d), hide_index=True)
        else:
            st.info("No reagents were found in this document.")
        if d["equipment"]:
            st.markdown("**Equipment**")
            st.markdown("\n".join(f"- {md_escape(x)}" for x in d["equipment"]))

    with t_safe:
        if d["safety_precautions"]:
            _stretch(st.dataframe, safety_df(d), hide_index=True)
        else:
            st.info("No safety precautions were found in this document.")
        st.warning(DISCLAIMER)

    with t_exp:
        base = safe_name(rec["filename"])
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            _stretch(
                st.download_button, "Word (.docx)", to_docx(d, rec["filename"]), f"{base}_summary.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"docx_{key}",
            )
        with c2:
            _stretch(st.download_button, "Markdown", to_markdown(d, rec["filename"]), f"{base}_summary.md", "text/markdown", key=f"md_{key}")
        with c3:
            _stretch(
                st.download_button, "Reagents (.csv)", reagents_df(d).to_csv(index=False).encode("utf-8-sig"),
                f"{base}_reagents.csv", "text/csv", key=f"csv_{key}",
            )
        with c4:
            _stretch(
                st.download_button, "JSON", json.dumps(d, indent=2, ensure_ascii=False), f"{base}_summary.json",
                "application/json", key=f"json_{key}",
            )


def run_analysis(client: Groq, ws: str, files, links) -> list[str]:
    ids: list[str] = []
    errors: list[str] = []
    items: list[tuple[str, str, bytes]] = []

    with st.status("Analyzing papers", expanded=True) as status:
        for f in files:
            if f.size > MAX_PDF_MB * 1024 * 1024:
                errors.append(f"{f.name}: larger than {MAX_PDF_MB} MB.")
                continue
            items.append((f.name, "Upload", f.getvalue()))
        for url in links:
            st.write("Downloading from Google Drive")
            try:
                name, content = fetch_drive_pdf(url)
                items.append((name, "Google Drive", content))
            except Exception as e:
                errors.append(f"{url}: {e}")

        progress = st.progress(0.0)
        for n, (name, source, raw) in enumerate(items, 1):
            try:
                ids.append(process_paper(client, ws, name, source, raw, st.write))
            except Exception as e:
                errors.append(f"{name}: {e}")
            progress.progress(n / max(len(items), 1))

        if ids:
            status.update(label=f"Done. {len(ids)} of {len(files) + len(links)} papers analysed.", state="complete", expanded=False)
        else:
            status.update(label="No papers could be analysed", state="error", expanded=True)

    for msg in errors:
        st.error(msg)
    return ids


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main() -> None:
    st.markdown(CSS, unsafe_allow_html=True)
    ws = get_workspace()
    purge_expired()
    render_header()
    files, links, analyze = render_sidebar()

    if "batch_ids" not in st.session_state:
        st.session_state.batch_ids = []

    if analyze:
        client = get_client()
        if client is None:
            st.error("The analysis service is not configured yet. The app owner needs to add GROQ_API_KEY to the app secrets.")
        else:
            st.session_state.batch_ids = run_analysis(client, ws, files, links)

    saved = list_results(ws)
    by_id = {r["id"]: r for r in saved}
    current = [by_id[i] for i in st.session_state.batch_ids if i in by_id]

    tab_current, tab_saved = st.tabs(["Current analysis", f"Saved for {RETENTION_DAYS} days ({len(saved)})"])

    with tab_current:
        if not current:
            st.markdown(
                """<div class="lab-empty">
                <b>Start with a paper.</b><br>
                Upload up to 5 PDFs or paste Google Drive links in the side panel, then select <b>Analyze papers</b>.
                You will get the experimental procedures, a reagent list and safety precautions for each paper.
                </div>""",
                unsafe_allow_html=True,
            )
        elif len(current) == 1:
            render_result(current[0], key=f"cur_{current[0]['id']}")
        else:
            labels = [f"{i}. {r['data']['title'] or r['filename']}" for i, r in enumerate(current, 1)]
            for tab, rec in zip(st.tabs([l[:32] for l in labels]), current):
                with tab:
                    render_result(rec, key=f"cur_{rec['id']}")
            combined = "\n\n---\n\n".join(to_markdown(r["data"], r["filename"]) for r in current)
            st.download_button("Download all as Markdown", combined, "lab_summaries.md", "text/markdown", key="all_md")

    with tab_saved:
        if not saved:
            st.info("Nothing saved yet. Your analyses will appear here for 7 days.")
        else:
            ids = [r["id"] for r in saved]
            choice = st.selectbox(
                "Saved analyses",
                ids,
                format_func=lambda i: f"{by_id[i]['data']['title'] or by_id[i]['filename']}  (saved {fmt_date(by_id[i]['created_at'])}, expires {fmt_date(by_id[i]['expires_at'])})",
            )
            render_result(by_id[choice], key=f"saved_{choice}")
            if st.button("Delete this analysis", key=f"del_{choice}"):
                delete_result(ws, choice)
                st.session_state.batch_ids = [i for i in st.session_state.batch_ids if i != choice]
                st.rerun()


main()
